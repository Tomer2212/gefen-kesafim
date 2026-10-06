import re
from pathlib import Path

import openpyxl
import pandas as pd

from logic.gefen_processor import normalize_amount

_DATE_SHAPE_RE = re.compile(r"^\d{2}\.\d{2}\.\d{2}$")
_PLAIN_INT_RE = re.compile(r"^\d+$")
_AMOUNT_SHAPE_RE = re.compile(r"^-?\d+(,\d{3})*(\.\d{1,2})?-?$")


def load_kesafim(filepath: str) -> pd.DataFrame:
    line_rows = _read_rows(filepath)
    rows = _parse_blocks(line_rows)
    if not rows:
        return pd.DataFrame(columns=[
            "report_code", "supplier", "supplier_name",
            "invoice_date", "invoice_number", "voucher",
            "item_number", "item_name", "description",
            "amount_raw", "total", "status", "amount", "ichud",
        ])
    df = pd.DataFrame(rows)
    df["amount"] = df["amount_raw"].apply(normalize_amount)
    df["ichud"] = (
        df["supplier"].astype(str)
        + "-"
        + df["invoice_number"].astype(str)
        + "-"
        + df["report_code"].astype(str)
        + "-"
        + df["amount"]
    )
    return df


# Legacy fixed column layout (old 13-column Kesafim2000 export). Used as a
# fallback for any field whose header text is not found in a block's header row.
_LEGACY_IDX = {
    "supplier": 0,
    "supplier_name": 1,
    "invoice_date": 2,
    "invoice_number": 3,
    "voucher": 4,
    "item_number": 5,
    "item_name": 6,
    "description": 7,
    "amount_raw": 10,
    "total": 11,
    "status": 12,
}

# Internal field name -> Hebrew column header as it appears in the row directly
# under each "קוד גפן" block header. The header text is identical between the old
# 13-column layout and the new 16-column layout (Kesafim2000 added 3 columns and
# reordered the rest) — only the position changed. Mapping by header name keeps
# both layouts working.
_HEADER_TO_FIELD = {
    "ספק": "supplier",
    "שם": "supplier_name",
    "תאריך חשבונית": "invoice_date",
    "מס.חשבונית": "invoice_number",
    "שובר הוצאה": "voucher",
    "מס פריט": "item_number",
    "שם פריט": "item_name",
    "מהות החשבונית": "description",
    "סכום פריט": "amount_raw",
    'סה"כ לחשבונית': "total",
    "סטטוס חשבונית": "status",
}


def _read_rows(filepath: str) -> list[list[str]]:
    """Return the file's rows as lists of stripped strings, regardless of whether
    the file is the classic Kesafim2000 export (.xls, actually TSV text in
    iso-8859-8) or a genuine .xlsx the file was re-saved as (e.g. to make it
    easier to locate/save). Both encode the same block structure (see
    _parse_blocks), just with different cell types — a real xlsx can hand back
    numbers (int/float) instead of text, so every cell is stringified here to
    keep that structure identical for both sources."""
    if Path(filepath).suffix.lower() == ".xlsx":
        rows = []
        wb = openpyxl.load_workbook(filepath, read_only=True)
        try:
            for sheet_name in wb.sheetnames:
                for row in wb[sheet_name].iter_rows(values_only=True):
                    rows.append(["" if v is None else str(v).strip() for v in row])
        finally:
            wb.close()
        return rows

    with open(filepath, "r", encoding="iso-8859-8") as f:
        content = f.read()
    return [line.rstrip("\r").split("\t") for line in content.strip().split("\n")]


def _looks_like_date(s: str) -> bool:
    return bool(_DATE_SHAPE_RE.match(s.strip()))


def _looks_like_plain_number(s: str) -> bool:
    return bool(_PLAIN_INT_RE.match(s.strip()))


def _looks_like_amount(s: str) -> bool:
    s = s.strip()
    if not _AMOUNT_SHAPE_RE.match(s):
        return False
    return normalize_amount(s) != ""


def _looks_like_invoice_number(s: str) -> bool:
    s = s.strip()
    return bool(s) and any(ch.isdigit() for ch in s)


def _try_splice_continuation(parts: list[str], line_rows: list[list[str]], idx: int, col_idx: dict) -> list[str] | None:
    """Recover a row broken across two physical lines by a stray line-break in
    the source export (observed real-world Kesafim2000 artifact: columns A,B
    land on one line, columns C.. land on the next with an empty leading cell).
    Deliberately scoped to the exact observed shape (only columns A,B present)
    rather than any truncation width. Returns the merged row, or None if this
    isn't that pattern (caller then falls back to today's existing drop)."""
    if len(parts) != 2 or not parts[0].strip() or not parts[1].strip():
        return None
    if idx + 1 >= len(line_rows):
        return None

    next_parts = line_rows[idx + 1]
    next_col_a = next_parts[0].strip() if next_parts else ""
    if next_col_a and _PLAIN_INT_RE.match(next_col_a):
        return None  # next line is its own legitimate data row

    merged = parts + next_parts[1:]

    def _merged_cell(field: str) -> str:
        i = col_idx.get(field)
        return merged[i].strip() if i is not None and i < len(merged) else ""

    if not _looks_like_date(_merged_cell("invoice_date")):
        return None
    if not _looks_like_plain_number(_merged_cell("item_number")):
        return None
    if not _looks_like_amount(_merged_cell("amount_raw")):
        return None
    if not _looks_like_amount(_merged_cell("total")):
        return None
    if not _looks_like_invoice_number(_merged_cell("invoice_number")):
        return None

    return merged


def _parse_blocks(line_rows: list[list[str]]) -> list[dict]:
    rows = []
    current_code = None
    header_next = False
    col_idx = dict(_LEGACY_IDX)  # field -> column index for the current block
    skip_next = False

    for idx, parts in enumerate(line_rows):
        if skip_next:
            skip_next = False
            continue
        if parts[0] == "קוד גפן":
            current_code = int(parts[1]) if parts[1].strip().isdigit() else None
            header_next = True
            continue
        if header_next:
            header_next = False
            # Rebuild the field -> index map from this block's header row.
            # Any field whose header is missing keeps its legacy fixed index.
            mapping = dict(_LEGACY_IDX)
            for i, cell in enumerate(parts):
                field = _HEADER_TO_FIELD.get(cell.strip())
                if field is not None:
                    mapping[field] = i
            col_idx = mapping
            continue
        if not parts[0].strip() or parts[0].strip() == " ":
            continue
        if current_code is None:
            continue

        # A valid data row must at least reach the amount column. If not, try
        # to recover a row that got split across two physical lines (see
        # _try_splice_continuation) before giving up on it as before.
        if col_idx["amount_raw"] >= len(parts):
            merged = _try_splice_continuation(parts, line_rows, idx, col_idx)
            if merged is None:
                continue
            parts = merged
            skip_next = True

        def _cell(field: str) -> str:
            i = col_idx.get(field)
            return parts[i].strip() if i is not None and i < len(parts) else ""

        rows.append({
            "report_code": current_code,
            "supplier": _cell("supplier"),
            "supplier_name": _cell("supplier_name"),
            "invoice_date": _cell("invoice_date"),
            "invoice_number": _cell("invoice_number"),
            "voucher": _cell("voucher"),
            "item_number": _cell("item_number"),
            "item_name": _cell("item_name"),
            "description": _cell("description"),
            "amount_raw": _cell("amount_raw"),
            "total": _cell("total"),
            "status": _cell("status"),
        })

    return rows
