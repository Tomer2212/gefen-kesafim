"""Combined PDF and Excel export across multiple sections."""

from io import BytesIO

import openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, SimpleDocTemplate

from logic.pdf_exporter import (
    _ensure_fonts,
    build_hashva_section_story,
    build_school_info_story,
)
from logic.tikhnun_exporter import (
    _write_kvua,
    _write_nihul,
    _write_partial,
    _write_sikar,
    _write_yozma,
    build_tikhnun_section_story,
)

TIKHNUN_TABS = {"sikar", "kvua", "partial", "yozma", "nihul"}
HASHVA_TABS = {"hashva", "rejected", "nopdf"}

_SHEET_NAMES = {
    "sikar":   "סקירה - גפן",
    "kvua":    "מימוש תקציב קבוע",
    "partial": "תוכניות עם ביצוע חלקי",
    "yozma":   "יוזמות וצרכים",
    "nihul":   "ניהול ותפעול",
}


def export_combined_pdf(run_data: dict, sections: list, multiplier: str = "03") -> bytes:
    _ensure_fonts()
    tikhnun = (run_data or {}).get("tikhnun")
    has_tikhnun = bool(tikhnun and not tikhnun.get("error"))
    stage_override = (run_data or {}).get("stage_override")

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        rightMargin=2*cm, leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm,
    )
    story = []

    if has_tikhnun:
        story.extend(build_school_info_story(tikhnun, stage_override))
        if sections:
            story.append(PageBreak())

    for i, section in enumerate(sections):
        if i > 0:
            story.append(PageBreak())
        if section in TIKHNUN_TABS and has_tikhnun:
            story.extend(build_tikhnun_section_story(tikhnun, section, multiplier, stage_override))
        elif section in HASHVA_TABS:
            story.extend(build_hashva_section_story(run_data or {}, section))

    if not story:
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import Paragraph
        story = [Paragraph("אין נתונים", ParagraphStyle("empty"))]

    doc.build(story)
    return buf.getvalue()


def export_combined_excel(run_data: dict, sections: list, multiplier: str = "03") -> bytes:
    tikhnun = (run_data or {}).get("tikhnun")
    has_tikhnun = bool(tikhnun and not tikhnun.get("error"))
    stage_override = (run_data or {}).get("stage_override")

    hashva_sections = [s for s in sections if s in HASHVA_TABS]
    tikhnun_sections = [s for s in sections if s in TIKHNUN_TABS]
    selected_budget = (run_data or {}).get("selected_budget_name")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    if hashva_sections:
        # Build every requested hashva-type sheet directly from run_data's row lists
        # (already budget-scoped upstream when relevant — see
        # analyze_router._scope_run_to_budget_combo), reusing excel_exporter's existing
        # sheet-writing helper as-is. _add_result_sheet already shows "✓ אין פערים" for
        # an empty list instead of omitting the sheet, so a clean budget/section still
        # gets its own tab. (There used to be a `run_data["file_path"]` branch here that
        # copied sheets from a pre-generated workbook — that key was never actually set
        # anywhere in the backend, so it silently never ran; this replaces it.)
        from logic.excel_exporter import _add_result_sheet, red_fill
        import pandas as pd
        finance_sw = (run_data.get("summary") or {}).get("finance_file", {}).get("software") or "כספים"
        budget_suffix = f"-{selected_budget}" if selected_budget else ""
        # Excel sheet titles are hard-capped at 31 chars — truncate explicitly.
        if "hashva" in hashva_sections:
            _add_result_sheet(
                wb, f"קיים ב{finance_sw} לא בגפן{budget_suffix}"[:31],
                pd.DataFrame(run_data.get("rows_finance_not_gefen") or []), red_fill,
            )
            _add_result_sheet(
                wb, f"משויך בגפן לא ב{finance_sw}{budget_suffix}"[:31],
                pd.DataFrame(run_data.get("rows_gefen_not_finance") or []), red_fill,
            )
        if "rejected" in hashva_sections:
            _add_result_sheet(
                wb, "אסמכתאות שנדחו",
                pd.DataFrame(run_data.get("rows_gefen_rejected") or []), red_fill,
            )
        if "nopdf" in hashva_sections:
            _add_result_sheet(
                wb, "אסמכתאות ללא PDF",
                pd.DataFrame(run_data.get("rows_gefen_no_pdf") or []), red_fill,
            )

    if has_tikhnun and tikhnun_sections:
        yozma_key = "yozma_04" if multiplier == "04" else "yozma_03"
        write_fns = {"sikar": _write_sikar, "kvua": _write_kvua, "partial": _write_partial}
        for section in tikhnun_sections:
            if section == "sikar":
                base_title = f"סקירה - {tikhnun.get('selected_budget_name') or 'גפן'}"[:31]
            else:
                base_title = _SHEET_NAMES.get(section, section)
            title = base_title
            n = 1
            while title in wb.sheetnames:
                n += 1
                title = f"{base_title} ({n})"
            ws = wb.create_sheet(title=title)
            ws.sheet_view.rightToLeft = True
            if section == "yozma":
                _write_yozma(ws, tikhnun, yozma_key)
            elif section == "nihul":
                _write_nihul(ws, tikhnun)
            elif section in ("sikar", "kvua"):
                write_fns[section](ws, tikhnun, stage_override)
            elif section in write_fns:
                write_fns[section](ws, tikhnun)

    if not wb.sheetnames:
        wb.create_sheet("נתונים")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
